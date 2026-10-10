"""Build an intake workbook and fill it with a fictional cohort of 59 students.

  python tests/make_example.py [examples/example_intake.xlsx]
"""
import random
import subprocess
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import layout as lay  # noqa: E402

KEY = {k: text for k, text, _ in lay.KEYS}
FORMS_HEADERS = (["Id", "Start time", "Completion time", KEY["email"], KEY["name"], "Last modified time",
                  KEY["discipline"]] + [f"{KEY['skills']}.{stmt}" for _, stmt in lay.SKILLS]
                 + [KEY["availability"], KEY["gender"], KEY["apart"]])


def build_template(path: Path) -> None:
    subprocess.run([sys.executable, str(ROOT / "build_workbook.py"), str(path)],
                   check=True, capture_output=True)


def email_of(name: str) -> str:
    return name.lower().replace(" ", ".") + "@example.edu"


def forms_rows(responses: list[dict]) -> list[list]:
    """Rows as a Microsoft Forms export would give them. Each dict: name and optionally email,
    completed (datetime), discipline, skills (9 ints), blocks (labels), gender, apart."""
    rows = [FORMS_HEADERS]
    start = datetime(2026, 10, 9, 14, 0)
    for n, r in enumerate(responses, 1):
        done = r.get("completed", start + timedelta(minutes=7 * n))
        skills = r.get("skills", [3] * len(lay.SKILLS))
        rows.append([n, (done - timedelta(minutes=4)).strftime("%m/%d/%y %H:%M:%S"),
                     done.strftime("%m/%d/%y %H:%M:%S"), r.get("email", email_of(r["name"])), r["name"], "",
                     r.get("discipline", lay.DISCIPLINES[0])]
                    + [lay.RATING_OPTIONS[v - 1] for v in skills]
                    + [";".join(r.get("blocks", lay.BLOCKS)) + ";", r.get("gender", ""), r.get("apart", "")])
    return rows


def fill(path: Path, roster: list, responses: list[dict] | None = None, overrides=(),
         settings: dict | None = None, rows: list[list] | None = None) -> None:
    """Write a roster (names, or (name, email) pairs), Forms style responses, override rows and Settings cells."""
    wb = load_workbook(path)
    for ref, value in (settings or {}).items():
        wb[lay.SETTINGS][ref] = value
    ro = wb[lay.ROSTER]
    for i, entry in enumerate(roster):
        name, email = (entry, email_of(entry)) if isinstance(entry, str) else entry
        ro.cell(lay.DATA_FIRST + i, 1, name)
        ro.cell(lay.DATA_FIRST + i, 2, email)
    rs = wb[lay.RESPONSES]
    rs.delete_rows(1, rs.max_row)
    for r, row in enumerate(rows if rows is not None else forms_rows(responses or []), 1):
        for j, v in enumerate(row, 1):
            rs.cell(r, j, v)
    ov = wb[lay.OVERRIDES]
    for i, row in enumerate(overrides):
        for j, v in enumerate(row, 1):
            ov.cell(lay.DATA_FIRST + i, j, v)
    wb.save(path)


PROFILES = {  # discipline: skills that tend to be high (indices into lay.SKILLS)
    "Psychology / cognitive science": [0, 6, 8],
    "Computer science / engineering": [3, 4, 5],
    "Design / architecture / arts": [1, 2, 7],
    "Business / management": [0, 8],
    "Social sciences / humanities": [0, 8, 6],
    "Natural sciences / medicine": [6, 4],
    "Other": [7],
}


def example_cohort(seed: int = 26):
    rng = random.Random(seed)
    first = ["Ada", "Ben", "Chloe", "Dev", "Emil", "Fatima", "Gus", "Hana", "Ivo", "Jun", "Kemi", "Lior",
             "Maya", "Nico", "Omar", "Priya", "Quinn", "Rosa", "Sami", "Tara", "Uma", "Vik", "Wen", "Xavi",
             "Yara", "Zed", "Alba", "Bo", "Cyra", "Dara"]
    names = [f"{f} {s}" for s in ("Example", "Sample") for f in first][:59]
    discs = (["Psychology / cognitive science"] * 18 + ["Computer science / engineering"] * 14
             + ["Design / architecture / arts"] * 10 + ["Business / management"] * 5
             + ["Social sciences / humanities"] * 6 + ["Natural sciences / medicine"] * 3 + ["Other"] * 3)
    genders = (["Woman"] * 33 + ["Man"] * 18 + [lay.GENDERS[2][0]] * 2 + ["Prefer not to say"] * 2)
    rng.shuffle(discs)
    rng.shuffle(genders)
    evening = [b for b in lay.BLOCKS if "evening" in b] + ["Saturday", "Sunday"]
    responders = names[:55]                       # the last 4 on the roster never answer
    responses = []
    for i, name in enumerate(responders):
        d = discs[i]
        skills = [min(5, max(1, rng.choice([1, 2, 2, 3]) + (2 if k in PROFILES[d] else 0)))
                  for k in range(len(lay.SKILLS))]
        if i % 11 == 5:                           # five part time students: evenings and weekends
            blocks = [b for b in evening if rng.random() < 0.85]
        else:
            blocks = [b for b in lay.BLOCKS if rng.random() < 0.65]
        responses.append(dict(name=name, discipline=d, skills=skills, blocks=blocks or ["Sunday"],
                              gender=genders[i]))
    responses[3]["apart"] = "Gus Exmple"           # misspelt, matched by fuzzy name matching
    responses[7]["apart"] = "Someone Unknown"      # no such student, listed in the report
    resubmit = dict(responses[10], completed=datetime(2026, 10, 10, 9, 0))
    responses[10] = dict(responses[10], blocks=["Sunday"])   # superseded by the later submission
    responses.append(resubmit)
    responses.append(dict(name="Wren Withdrawn", discipline="Other"))   # not on the roster
    overrides = [("Hana Example", "Team 03", "", "Reasonable adjustment (fictional)"),
                 ("Ivo Example", "", "Jun Example", "Staff request (fictional)")]
    return names, responses, overrides


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else ROOT / "examples" / "example_intake.xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    build_template(target)
    roster, responses, overrides = example_cohort()
    fill(target, roster, responses, overrides)
    print("wrote", target)
