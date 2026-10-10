"""Run with:  python -m pytest -q   (after: pip install -r requirements-dev.txt)"""
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import allocate_teams as at  # noqa: E402
import layout as lay  # noqa: E402
from make_example import build_template, email_of, example_cohort, fill, forms_rows  # noqa: E402

WEAK, STRONG = [2] * 9, [5] * 9
ALL = lay.BLOCKS


@pytest.fixture(scope="session")
def template(tmp_path_factory):
    path = tmp_path_factory.mktemp("tpl") / "template.xlsx"
    build_template(path)
    return path.read_bytes()


def workbook(tmp_path, template, roster, responses, overrides=(), settings=None, rows=None):
    path = tmp_path / "intake.xlsx"
    path.write_bytes(template)
    fill(path, roster, responses, overrides, settings, rows)
    return path


def run(path, responses=None, time_limit=20):
    p = at.read_problem(path, responses, None)
    res = at.solve(p, at.DEFAULT_PRIORITY, time_limit, True, 0)
    return p, res


def team_of(p):
    return {s.name: s.team for s in p.students}


def people(names, **kw):
    return [dict(name=n, **kw) for n in names]


def test_example_end_to_end(tmp_path, template):
    roster, responses, overrides = example_cohort()
    path = workbook(tmp_path, template, roster, responses, overrides)
    assert at.main([str(path), "--time-limit", "20"]) == 0
    out = load_workbook(tmp_path / "intake_teams.xlsx")
    assert out.sheetnames == ["Report", "Teams", "Students"]
    assert (tmp_path / "intake_teams.csv").exists()

    p, res = run(path, time_limit=10)
    sizes = Counter(s.team for s in p.students)
    assert len(p.students) == 59 and sorted(sizes.values()) == [4] + [5] * 11
    assert {s.name for s in p.students} == set(roster)                 # roster is authoritative
    assert [who for who, _ in p.excluded] == ["Wren Withdrawn <wren.withdrawn@example.edu>"]
    t = team_of(p)
    assert t["Hana Example"] == 2                                       # pinned to Team 03
    assert t["Ivo Example"] != t["Jun Example"]                         # staff keep apart
    assert t["Dev Example"] != t["Gus Example"]                         # student request, misspelt
    assert ("Hana Example", "Someone Unknown") in p.unmatched
    kemi = next(s for s in p.students if s.name == "Kemi Example")
    assert len(kemi.blocks) > 1                                         # later submission kept
    # independent checks of the evidence based criteria
    for g in range(p.n_teams):
        mem = [s for s in p.students if s.team == g]
        for gender in ("woman", "man"):
            assert sum(s.gender == gender for s in mem) != 1
        known = [s.blocks for s in mem if s.blocks is not None]
        assert len(set.intersection(*known)) >= 2
        assert sum(not s.responded for s in mem) <= 1
    assert res.terms["lone"] == 0 and res.terms["time"] == 0


def test_lone_member_rule_outranks_skills(tmp_path, template):
    names = [f"S{i}" for i in range(10)]
    resp = people(names[:8], skills=WEAK, gender="Man")
    resp += [dict(name="W1", skills=[5, 5, 5, 5, 1, 1, 1, 1, 1], gender="Woman"),
             dict(name="W2", skills=[5, 5, 5, 5, 1, 1, 1, 1, 1], gender="Woman")]
    roster = names[:8] + ["W1", "W2"]
    p, _ = run(workbook(tmp_path, template, roster, resp))
    t = team_of(p)
    assert t["W1"] == t["W2"]          # separating them would cover more skills, but isolate both


def test_keep_apart_is_hard(tmp_path, template):
    roster = [f"S{i}" for i in range(8)] + ["W1", "W2"]
    resp = people(roster[:8], skills=WEAK, gender="Man")
    resp += [dict(name="W1", gender="Woman", apart="W2"), dict(name="W2", gender="Woman")]
    p, res = run(workbook(tmp_path, template, roster, resp))
    t = team_of(p)
    assert t["W1"] != t["W2"] and res.terms["lone"] == 2   # apart wins over the lone member rule


def test_shared_free_time_outranks_skills(tmp_path, template):
    mon, wed = ["Monday morning", "Monday evening"], ["Wednesday morning", "Wednesday afternoon"]
    a = [dict(name=f"A{i}", skills=[5] * 5 + [1] * 4, blocks=mon) for i in range(5)]
    b = [dict(name=f"B{i}", skills=[1] * 5 + [5] * 4, blocks=wed) for i in range(5)]
    roster = [s["name"] for s in a + b]
    p, res = run(workbook(tmp_path, template, roster, a + b))
    t = team_of(p)
    assert len({t[f"A{i}"] for i in range(5)}) == 1 and len({t[f"B{i}"] for i in range(5)}) == 1
    assert res.terms["time"] == 0


def test_time_outranks_any_amount_of_skill_coverage(tmp_path, template):
    """The weighted model of version 2.0 gave up a shared block to close nine skill gaps; strict priority does not."""
    mon, wed = ["Monday morning", "Monday evening"], ["Wednesday morning", "Wednesday afternoon"]
    resp = [dict(name=f"A{i}", skills=STRONG, blocks=mon) for i in range(4)]
    resp += [dict(name="P", skills=STRONG, blocks=mon + ["Wednesday morning"])]
    resp += [dict(name=f"B{i}", skills=[1] * 9, blocks=wed) for i in range(4)]
    resp += [dict(name="X", skills=[1] * 9, blocks=ALL)]
    p, res = run(workbook(tmp_path, template, [r["name"] for r in resp], resp))
    t = team_of(p)
    assert res.terms["time"] == 0 and t["P"] == t["A0"] and t["X"] == t["B0"]


def test_non_responders_are_spread(tmp_path, template):
    roster = [f"S{i:02d}" for i in range(15)]
    p, res = run(workbook(tmp_path, template, roster, people(roster[:12], skills=WEAK)))
    per_team = Counter(s.team for s in p.students if not s.responded)
    assert len(p.students) == 15 and sorted(per_team.values()) == [1, 1, 1]


def test_forms_parsing_and_latest_submission(tmp_path, template):
    roster = [f"S{i}" for i in range(5)]
    resp = people(roster, skills=WEAK)
    resp.append(dict(name="S0", skills=STRONG, completed=datetime(2026, 10, 11, 10, 0)))
    resp.insert(0, dict(name="S1", skills=STRONG, completed=datetime(2026, 10, 12, 10, 0)))
    rows = forms_rows(resp)
    rows[1][-3] = "Monday morning;Moon base;"             # unknown block label
    rows.insert(1, [None] * len(rows[0]))                 # a blank row: the label is now in sheet row 3
    p, _ = run(workbook(tmp_path, template, roster, None, rows=rows))
    s = {st.name: st for st in p.students}
    assert s["S0"].skills == STRONG                        # later submission, later row
    assert s["S1"].skills == STRONG                        # later submission, earlier row
    assert any("row 3" in w and "Moon base" in w for w in p.warnings)
    assert "replaced by the latest" in p.notes[0]


def test_separate_csv_export_and_missing_gender_column(tmp_path, template):
    import csv
    roster = [f"S{i}" for i in range(6)]
    rows = forms_rows(people(roster, skills=WEAK))
    gender = rows[0].index(lay.KEYS[6][1])
    rows = [r[:gender] + r[gender + 1:] for r in rows]
    export = tmp_path / "forms.csv"
    with open(export, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(rows)
    path = workbook(tmp_path, template, roster, [])
    p, _ = run(path, responses=export)
    assert all(s.responded for s in p.students)
    assert any("gender" in w for w in p.warnings)


def test_forms_xlsx_export_with_stale_dimension(tmp_path, template):
    """A real Forms export declares <dimension ref="A1"/>; read-only openpyxl then sees only column A."""
    import re
    import zipfile
    from openpyxl import Workbook
    roster = [f"S{i}" for i in range(6)]
    wb = Workbook()
    for row in forms_rows(people(roster, skills=WEAK)):
        wb.active.append(row)
    plain = tmp_path / "plain.xlsx"
    wb.save(plain)
    export = tmp_path / "forms.xlsx"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(export, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item.startswith("xl/worksheets/"):
                data = re.sub(rb"<dimension[^>]*/>", b'<dimension ref="A1" />', data)
            zout.writestr(item, data)
    path = workbook(tmp_path, template, roster, [])
    p, _ = run(path, responses=export)
    assert all(s.responded for s in p.students)


def test_google_forms_csv_export(tmp_path, template):
    """Google Forms: 'Timestamp', 'Email Address', grid headers 'Question [Statement]', checkboxes joined by ', '."""
    import csv
    roster = [f"S{i}" for i in range(6)]
    rows = [["Timestamp", "Email Address", lay.KEYS[3][1]] + [f"{lay.KEYS[4][1]} [{stmt}]" for _, stmt in lay.SKILLS]
            + [lay.KEYS[5][1], lay.KEYS[6][1]]]
    for i, name in enumerate(roster):
        rows.append([f"2026/10/0{i + 1} 2:27:07 PM GMT+1", email_of(name), lay.DISCIPLINES[0]] + ["4"] * len(lay.SKILLS)
                    + ["Monday morning, Wednesday afternoon", "Woman" if i % 2 else "Man"])
    export = tmp_path / "google.csv"
    with open(export, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    p, _ = run(workbook(tmp_path, template, roster, []), responses=export)
    assert all(s.responded for s in p.students) and not p.excluded
    both = {lay.BLOCKS.index("Monday morning"), lay.BLOCKS.index("Wednesday afternoon")}
    assert all(s.blocks == both and s.skills == [4] * len(lay.SKILLS) for s in p.students)


def test_responses_linked_by_name_in_any_word_order(tmp_path, template):
    """Microsoft Forms gives the account name ('Lin Mei', 'Lind, Eva') and a username, not the roster email."""
    roster = ["Mei Lin", "Eva Lind", "Ana-Maria Costa", "Ola Berg", "Alex Gray", "Sam Lee",
              "Ria D'Souza", "Jo-Anne Smith"]
    asked = ["Lin Mei", "Lind, Eva", "Costa, Ana-Maria", "Berg Ola", "Alex Gray", "Lee Sam",
             "D’Souza, Ria", "Jo Anne Smith"]
    responses = [dict(name=n, email=f"user{i:03d}@example.edu", skills=WEAK) for i, n in enumerate(asked)]
    path = workbook(tmp_path, template, roster, responses)
    p, _ = run(path)
    assert not p.excluded
    assert all(s.responded for s in p.students)
    assert sum(what.startswith("Linked by a differently written name") for what, _ in p.checks) == 7


def test_ambiguous_names_are_not_linked(tmp_path, template):
    """Two roster students with one name, or with swapped names, are never guessed between."""
    roster = [("Yu Wang", "yu.wang@example.edu"), ("Yu Wang", "yu.wang2@example.edu"), "Wei Li", "Li Wei",
              "Kai Chen", "Chen Kai", "S1", "S2"]
    resp = [dict(name="Yu Wang", email="user001@example.edu", apart="S1"), dict(name="Wang Yu", email="user002@example.edu"),
            dict(name="Li Wei", email="user003@example.edu"), dict(name="Kai, Chen", email="user004@example.edu")]
    p, _ = run(workbook(tmp_path, template, roster, resp))
    assert len(p.excluded) == 4 and not any(s.responded for s in p.students)
    assert all("fits 2 roster students" in why for _, why in p.excluded)
    assert "keep apart request 'S1' is not applied" in p.excluded[0][1]


def test_two_accounts_for_one_student_latest_counts(tmp_path, template):
    roster = ["Mei Lin", "S1", "S2", "S3"]
    resp = [dict(name="Mei Lin", email="user001@example.edu", skills=STRONG, completed=datetime(2026, 10, 10, 9, 0)),
            dict(name="Lin Mei", email="user002@example.edu", skills=WEAK, completed=datetime(2026, 10, 9, 9, 0))]
    p, _ = run(workbook(tmp_path, template, roster, resp + people(roster[1:], skills=WEAK)))
    assert next(s for s in p.students if s.name == "Mei Lin").skills == STRONG   # later, in the earlier row
    assert any("Two responses link to Mei Lin" in w for w in p.warnings)


def test_keep_apart_texts(tmp_path, template):
    roster = ["Alex Gray", "Sam Lee", "Lee Jones", "Ria D'Souza", "Ana-Maria Costa", "Bo Chen", "Cy Park", "Di Ross"]
    resp = people(roster, skills=WEAK)
    resp[0]["apart"] = "Lee, Sam"                        # one person, surname first: not Lee Jones
    resp[2]["apart"] = "D’Souza, Ria"              # typographic apostrophe
    resp[5]["apart"] = "N/A"
    resp[6]["apart"] = "Ana Maria Costa and Di Ross"
    p, _ = run(workbook(tmp_path, template, roster, resp))
    apart = {s.name: {p.students[j].name for j in s.apart} for s in p.students}
    assert apart["Alex Gray"] == {"Sam Lee"} and apart["Lee Jones"] == {"Ria D'Souza"}
    assert apart["Bo Chen"] == set() and apart["Cy Park"] == {"Ana-Maria Costa", "Di Ross"}
    assert p.unmatched == [] and len(p.apart_log) == 3


def test_gender_answers_outside_settings_are_reported(tmp_path, template):
    roster = [f"S{i}" for i in range(8)]
    resp = people(roster[:4], skills=WEAK, gender="Woman") + people(roster[4:6], skills=WEAK, gender="Man")
    resp += [dict(name="S6", gender="Non-binary"), dict(name="S7", gender="Female")]
    p, _ = run(workbook(tmp_path, template, roster, resp))
    assert any("'Non-binary' read as" in txt for _, txt in p.checks)
    assert any("'Female'" in w for w in p.warnings)


def test_report_numbers_are_measured_on_the_teams(tmp_path, template):
    roster, responses, _ = example_cohort()
    path = workbook(tmp_path, template, roster[:20], responses[:20])
    p = at.read_problem(path, None, None)
    res = at.solve(p, ["lone"], 5, True, 0)              # skill coverage is not optimised at all
    cfg = p.cfg
    cover = sum(any(s.skills[k] >= cfg.threshold for s in p.students if s.team == g)
                for g in range(p.n_teams) for k in range(len(cfg.skill_stmts)))
    assert res.terms["cover"] == cover > 0
    with pytest.raises(SystemExit):
        at.parse_args(["intake.xlsx", "--priority", "lone,bogus"])


def test_overrides_fuzzy_name_and_pin(tmp_path, template):
    roster = [f"Student {c}" for c in "ABCDEFGHIJ"]
    path = workbook(tmp_path, template, roster, people(roster, skills=WEAK),
                    overrides=[("Studnet A", "Team 02", "student b", "")])
    p, _ = run(path)
    t = team_of(p)
    assert t["Student A"] == 1 and t["Student B"] != 1


def test_too_few_teams_is_an_error(tmp_path, template):
    roster = [f"S{i}" for i in range(11)]
    path = workbook(tmp_path, template, roster, people(roster), settings={"B5": 2})
    with pytest.raises(SystemExit):
        at.read_problem(path, None, None)


def test_deterministic_runs_are_identical(tmp_path, template):
    roster, responses, _ = example_cohort()
    path = workbook(tmp_path, template, roster[:25], responses[:25])
    teams = [[s.team for s in run(path, time_limit=3)[0].students] for _ in range(2)]
    assert teams[0] == teams[1]
