#!/usr/bin/env python3
"""Form student project teams from the staff intake workbook.

Reads Settings, Roster, Responses and Overrides from the intake workbook (Responses can also
come from a separate Microsoft Forms or Google Forms export with --responses). Every student on the roster is
assigned to a team with Google OR-Tools CP-SAT. The input files are never modified.

Hard constraints
  every student is in exactly one team; team sizes differ by at most one;
  pinned students stay in their team; keep apart requests are always met.
Criteria in strict order of priority (change the order, or leave criteria out, with --priority)
  lone      fewest (team, gender) pairs with exactly one member of a gender marked Yes in Settings
  time      fewest shared free time blocks missing below the required number, over all teams
  worst     fewest skill gaps in the team with the most gaps (max-min fairness)
  nonresp   fewest non responders beyond the first in a team, over all teams
  cover     most (team, skill) pairs with a strong member (rating >= threshold)
  div       most distinct disciplines, summed over teams
  partial   most (team, skill) pairs with a member at threshold - 1 or above
Each criterion is optimised in turn and then held at its best value while the next one is
optimised (lexicographic optimisation), so a lower criterion never outweighs a higher one.

Outputs, written next to the workbook (or to --out-dir)
  <name>_teams.xlsx   Report, Teams and Students sheets
  <name>_teams.csv    Team, Mentor, Full name, Email (to create the groups in a learning platform)

Usage
  python allocate_teams.py team_formation.xlsx
  python allocate_teams.py intake.xlsx --responses forms_export.xlsx --deterministic --seed 1
"""
from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import math
import platform
import re
import sys
import unicodedata
import warnings
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from ortools import __version__ as ortools_version
from ortools.sat.python import cp_model

import layout as lay

__version__ = "2.1.0"

CRITERIA = {  # name: (minimise or maximise, Report label), in the default order of priority
    "lone": ("min", "Lone members ((team, gender) pairs, rule on)"),
    "time": ("min", "Missing shared free blocks (total)"),
    "worst": ("min", "Largest number of skill gaps in a team"),
    "nonresp": ("min", "Extra non responders in a team (total)"),
    "cover": ("max", "Strong (team, skill) pairs covered"),
    "div": ("max", "Distinct disciplines (sum over teams)"),
    "partial": ("max", "Partial (team, skill) pairs covered"),
}
DEFAULT_PRIORITY = list(CRITERIA)
NONE_WORDS = {"", "-", "none", "n/a", "na", "no", "nobody", "no one", "noone", "nope", "tbc", "?"}
APOSTROPHES = str.maketrans({c: "'" for c in "’‘ʼ`´"})


# ---------------------------------------------------------------------------- helpers
def norm(text) -> str:
    """Lower case ASCII with collapsed whitespace; typographic apostrophes become '."""
    text = unicodedata.normalize("NFKD", str(text or "").translate(APOSTROPHES)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().casefold()


def key(text) -> str:
    """Letters and digits only, for matching column headers and option labels."""
    return re.sub(r"[^a-z0-9]+", " ", norm(text)).strip()


def name_key(text) -> str:
    """Name words in alphabetical order, ignoring apostrophes, hyphens and commas, so 'Lind, Eva',
    'Costa, Ana Maria' and 'D’Souza, Ria' match 'Eva Lind', 'Ana-Maria Costa' and "Ria D'Souza"."""
    return " ".join(sorted(re.findall(r"[a-z0-9]+", norm(text).replace("'", ""))))


def is_none(text) -> bool:
    return norm(text).strip(" .") in NONE_WORDS


def as_str(v) -> str:
    return "" if v is None else str(v).strip()


def as_int(v) -> int | None:
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def rating(v) -> int:
    """First digit 1 to 5 in a cell such as 4, 4.0 or '3 done it in a project'; 0 if none."""
    m = re.search(r"(?<!\d)([1-5])(?!\d)", as_str(v))
    return int(m.group(1)) if m else 0


def is_example(*cells) -> bool:
    return any(norm(c).startswith(lay.EXAMPLE_PREFIX) for c in cells)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ---------------------------------------------------------------------------- data model
@dataclass
class Student:
    name: str
    email: str
    responded: bool = False
    discipline: str = ""
    skills: list[int] = field(default_factory=list)
    blocks: set[int] | None = None      # None: unknown, treated as free in every block
    gender: str = ""                    # key() of the Settings option
    apart_raw: str = ""
    row: int = 0                        # row of the response in the Responses sheet or export
    submitted: datetime | None = None
    pin: int | None = None              # team index
    apart: set[int] = field(default_factory=set)
    team: int | None = None


def newer(a: Student, b: Student) -> bool:
    """True if response a replaces response b: the later completion time, else the later row."""
    if a.submitted is not None and b.submitted is not None and a.submitted != b.submitted:
        return a.submitted > b.submitted
    return a.row > b.row


@dataclass
class Config:
    max_size: int
    min_size: int
    n_teams: int | None
    threshold: int
    req_blocks: int
    team_names: list[str]
    mentors: list[str]
    skill_labels: list[str]
    skill_stmts: list[str]
    disciplines: list[str]
    blocks: list[str]
    gender_rule: dict[str, bool]        # key(option) -> rule applies
    gender_label: dict[str, str]        # key(option) -> option as written
    tas: list[str]
    keys: dict[str, str]


@dataclass
class Problem:
    source: Path
    cfg: Config
    students: list[Student]
    n_teams: int
    warnings: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)                  # neutral information for the report
    mapping: list[tuple[str, str]] = field(default_factory=list)
    unmatched: list[tuple[str, str]] = field(default_factory=list)
    excluded: list[tuple[str, str]] = field(default_factory=list)   # responses not used: (who, why)
    checks: list[tuple[str, str]] = field(default_factory=list)     # matches staff should check
    apart_log: list[tuple[str, str]] = field(default_factory=list)  # keep apart requests as resolved
    groups: dict[str, list[int]] = field(default_factory=dict)      # lone member rule: gender -> students


# ---------------------------------------------------------------------------- reading
def read_settings(wb) -> Config:
    st = wb[lay.SETTINGS]

    def value(row, label, lo, hi, blank_ok=False):
        raw = st.cell(row, 2).value
        if blank_ok and as_str(raw) == "":
            return None
        v = as_int(raw)
        if v is None or not lo <= v <= hi:
            sys.exit(f"Settings!B{row} ({label}) must be a whole number from {lo} to {hi}, got {raw!r}.")
        return v

    def column(col, slots):
        return [as_str(st.cell(lay.TABLE_FIRST + i, col).value) for i in range(slots)]

    max_size = value(lay.MAX_SIZE, "maximum team size", 2, 8)
    min_size = value(lay.MIN_SIZE, "minimum team size", 1, 8)
    if min_size > max_size:
        sys.exit(f"Minimum team size ({min_size}) is larger than the maximum ({max_size}).")
    labels, stmts = column(lay.SKILL_LABEL_COL, lay.SKILL_SLOTS), column(lay.SKILL_STMT_COL, lay.SKILL_SLOTS)
    skills = [(lab or stm, stm) for lab, stm in zip(labels, stmts) if stm]
    if not skills:
        sys.exit("Settings has no skill statements (column E).")
    genders = column(lay.GENDER_COL, lay.GENDER_SLOTS)
    rules = column(lay.GENDER_RULE_COL, lay.GENDER_SLOTS)
    keys = {}
    for i, (k, default, _) in enumerate(lay.KEYS):
        keys[k] = as_str(st.cell(lay.TABLE_FIRST + i, lay.KEY_TEXT_COL).value) or default
    team_names = [n or f"Team {i + 1:02d}" for i, n in enumerate(column(lay.TEAM_COL, lay.TEAM_SLOTS))]
    if len({norm(n) for n in team_names}) != len(team_names):
        sys.exit("Team names in Settings must be unique.")
    return Config(
        max_size=max_size, min_size=min_size,
        n_teams=value(lay.N_TEAMS, "number of teams", 1, lay.TEAM_SLOTS, blank_ok=True),
        threshold=value(lay.THRESHOLD, "strong skill threshold", 1, 5),
        req_blocks=value(lay.REQ_BLOCKS, "required shared blocks", 0, lay.BLOCK_SLOTS),
        team_names=team_names, mentors=column(lay.MENTOR_COL, lay.TEAM_SLOTS),
        skill_labels=[a for a, _ in skills], skill_stmts=[b for _, b in skills],
        disciplines=[d for d in column(lay.DISC_COL, lay.DISC_SLOTS) if d],
        blocks=[b for b in column(lay.BLOCK_COL, lay.BLOCK_SLOTS) if b],
        gender_rule={key(g): norm(r) in ("yes", "y", "true", "1") for g, r in zip(genders, rules) if g},
        gender_label={key(g): g for g in genders if g},
        tas=[t for t in column(lay.TA_COL, lay.TA_SLOTS) if t], keys=keys)


def read_table(path: Path | None, wb) -> list[tuple[int, list]]:
    """Non-empty rows, with their row numbers, of the Responses sheet or of the first sheet / CSV
    given with --responses."""
    if path is None:
        rows = [list(r) for r in wb[lay.RESPONSES].iter_rows(values_only=True)]
    elif path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = [list(r) for r in csv.reader(f)]
    else:
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        ws.reset_dimensions()       # Forms exports declare the dimension as A1; read-only mode trusts it
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return [(n, r) for n, r in enumerate(rows, 1) if any(as_str(c) for c in r)]


def find_column(headers: list[str], text: str, contains_only: bool = False) -> int | None:
    """Exact header first, then a header starting with the text, then one containing it."""
    t = key(text)
    if not t:
        return None
    hk = [key(h) for h in headers]
    tests = [lambda h: t in h] if contains_only else \
        [lambda h: h == t, lambda h: h.startswith(t), lambda h: t in h]
    for test in tests:
        hits = [j for j, h in enumerate(hk) if test(h)]
        if len(hits) == 1:
            return hits[0]
        if len(hits) > 1:
            sys.exit(f"More than one Responses column matches '{text}': "
                     f"{', '.join(headers[j] for j in hits)}. Make the title in Settings more specific.")
    return None


def parse_time(v):
    if isinstance(v, datetime):
        return v
    for fmt in ("%m/%d/%y %H:%M:%S", "%m/%d/%y %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
                "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(as_str(v), fmt)
        except ValueError:
            pass
    return None


def read_responses(rows: list[tuple[int, list]], cfg: Config, p_warn: list[str], mapping: list[tuple[str, str]],
                   checks: list[tuple[str, str]]) -> tuple[list[Student], int]:
    if not rows:
        return [], 0
    headers = [as_str(h) for h in rows[0][1]]
    col = {k: find_column(headers, cfg.keys[k]) for k in
           ("name", "email", "completed", "discipline", "availability", "gender", "apart")}
    skill_cols = [find_column(headers, s, contains_only=True) for s in cfg.skill_stmts]
    if col["name"] is None and col["email"] is None:
        sys.exit("Responses: no 'name' or 'email' column found. Headers seen: " + "; ".join(headers))
    for k, label in (("name", "Name"), ("email", "Email"), ("completed", "Completion time"),
                     ("discipline", "Discipline"), ("availability", "Availability"),
                     ("gender", "Gender"), ("apart", "Keep apart")):
        found = headers[col[k]] if col[k] is not None else "NOT FOUND, criterion off"
        mapping.append((label, found))
        if col[k] is None and k == "email":
            p_warn.append("No Responses column for email: responses are matched to the roster by name, and two "
                          "responses with the same name count as one student's repeat submission.")
        elif col[k] is None and k in ("discipline", "availability", "gender", "apart"):
            p_warn.append(f"No Responses column for {label.lower()} ('{cfg.keys[k]}'). "
                          f"That criterion is off for this run.")
    for lab, j in zip(cfg.skill_labels, skill_cols):
        mapping.append((f"Skill: {lab}", headers[j] if j is not None else "NOT FOUND, rated 0"))
        if j is None:
            p_warn.append(f"No Responses column for the skill '{lab}'. Everyone is rated 0 on it.")

    block_index = {key(b): i for i, b in enumerate(cfg.blocks)}
    disc_lookup = {key(d): d for d in cfg.disciplines}
    read_as, unknown = Counter(), Counter()

    def gender_of(raw: str) -> str:
        """key() of the Settings option; an answer that starts like exactly one option counts as it."""
        k = key(raw)
        if not k or k in cfg.gender_label:
            return k
        hits = [o for o in cfg.gender_label if o.startswith(k + " ") or k.startswith(o + " ")]
        if len(hits) == 1:
            read_as[raw, cfg.gender_label[hits[0]]] += 1
            return hits[0]
        unknown[raw] += 1
        return k

    get = lambda r, k: as_str(r[col[k]]) if col[k] is not None and col[k] < len(r) else ""
    kept: dict[str, Student] = {}
    n_rows = 0
    for n, r in rows[1:]:
        name, email = get(r, "name"), get(r, "email")
        if not (name or email) or is_example(r[0] if r else "", name):
            continue
        n_rows += 1
        s = Student(name=name, email=email, responded=True, row=n)
        s.skills = [rating(r[j]) if j is not None and j < len(r) else 0 for j in skill_cols]
        raw_disc = get(r, "discipline")
        s.discipline = disc_lookup.get(key(raw_disc), raw_disc)
        if col["availability"] is not None:
            tokens = [t for t in re.split(r"[;,\n|]", get(r, "availability")) if key(t)]   # Forms ; Google ,
            chosen = {block_index[key(t)] for t in tokens if key(t) in block_index}
            unknown_blocks = [t.strip() for t in tokens if key(t) not in block_index]
            if unknown_blocks:
                p_warn.append(f"Responses row {n} ({name or email}): unknown time block(s) "
                              f"{', '.join(unknown_blocks)} ignored.")
            if chosen:
                s.blocks = chosen
            else:
                p_warn.append(f"Responses row {n} ({name or email}): no free time block given, "
                              f"treated as free in every block.")
        s.gender = gender_of(get(r, "gender"))
        s.apart_raw = get(r, "apart")
        if col["completed"] is not None and col["completed"] < len(r):
            s.submitted = parse_time(r[col["completed"]])
        ident = norm(email) or norm(name)
        prev = kept.get(ident)
        if prev is None or newer(s, prev):         # the latest submission counts
            kept[ident] = s
    for (raw, option), c in read_as.items():
        checks.append(("Gender answer read as a Settings option", f"'{raw}' read as '{option}' ({c} response(s))"))
    for raw, c in unknown.items():
        p_warn.append(f"Gender answer '{raw}' ({c} response(s)) is not an option in Settings, so no lone member "
                      f"rule applies to it. Add it to Settings column K if it should.")
    return sorted(kept.values(), key=lambda s: s.row), n_rows - len(kept)


def read_roster(wb) -> list[Student]:
    ws = wb[lay.ROSTER]
    out, seen = [], set()
    for r in range(lay.DATA_FIRST, lay.DATA_FIRST + lay.ROSTER_ROWS):
        name, email = as_str(ws.cell(r, 1).value), as_str(ws.cell(r, 2).value)
        if not (name or email) or is_example(name):
            continue
        ident = norm(email) or norm(name)
        if ident in seen:
            continue
        seen.add(ident)
        out.append(Student(name=name, email=email))
    return out


class Matcher:
    """Resolve a free text name or email to one student: email, name in any word order, close
    spelling, then a unique first name. Returns (index, how) or (None, why)."""

    def __init__(self, students: list[Student]):
        self.by_email = {norm(s.email): i for i, s in enumerate(students) if s.email}
        self.by_words: dict[str, list[int]] = {}
        self.by_first: dict[str, list[int]] = {}
        for i, s in enumerate(students):
            if name_key(s.name):
                self.by_words.setdefault(name_key(s.name), []).append(i)
                self.by_first.setdefault(name_key(norm(s.name).split(" ")[0]), []).append(i)

    def __call__(self, text: str) -> tuple[int | None, str]:
        t = norm(text).strip(" .")
        if t in self.by_email:
            return self.by_email[t], "email"
        k = name_key(t)
        hits = self.by_words.get(k, [])
        if not k or len(hits) > 1:
            return None, "fits more than one student" if k else "no name"
        if hits:
            return hits[0], "name"
        close = difflib.get_close_matches(k, list(self.by_words), n=1, cutoff=0.85)
        if close and len(self.by_words[close[0]]) == 1:
            return self.by_words[close[0]][0], "close spelling"
        if " " not in k and len(self.by_first.get(k, [])) == 1:
            return self.by_first[k][0], "first name only"
        return None, "no match"


def resolve_names(text: str, match: Matcher) -> tuple[list[tuple[int, str, str]], list[tuple[str, str]]]:
    """Students named in a keep apart text as (index, text, how), and the parts that matched nobody
    as (text, why). Each line or ';' part is first tried as one name, so 'Lee, Sam' stays one person;
    only if that fails is it split at commas, '/', '&', '+' and 'and'."""
    found, missed = [], []
    if is_none(text):
        return found, missed
    for chunk in re.split(r"[;\n]", text):
        if is_none(chunk):
            continue
        j, how = match(chunk)
        parts = [q for q in re.split(r"[,/&+]|\band\b", chunk, flags=re.IGNORECASE) if not is_none(q)]
        if j is not None:
            found.append((j, chunk.strip(), how))
        elif len(parts) < 2:
            missed.append((chunk.strip(), how))
        else:
            for q in parts:
                jq, hq = match(q)
                if jq is None:
                    missed.append((q.strip(), hq))
                else:
                    found.append((jq, q.strip(), hq))
    return found, missed


def apply_apart(p: Problem, match: Matcher, i: int, text: str, source: str) -> None:
    """Keep student i apart from everyone named in text; log how each name was matched."""
    found, missed = resolve_names(text, match)
    for tok, why in missed:
        p.unmatched.append((source, tok if why == "no match" else f"{tok} ({why})"))
    names = []
    for j, _, how in found:
        if j != i:
            p.students[i].apart.add(j)
            p.students[j].apart.add(i)
            names.append(p.students[j].name + ("" if how in ("email", "name") else f" ({how}, check)"))
    if found or missed:
        p.apart_log.append((source, f"'{text}' -> {', '.join(names) or 'nobody'}"))


def rule_groups(p: Problem) -> dict[str, list[int]]:
    """Gender options with the lone member rule on, mapped to their students."""
    groups = {}
    for g, on in p.cfg.gender_rule.items():
        if not on:
            continue
        members = [i for i, s in enumerate(p.students) if s.gender == g]
        if len(members) == 1:
            p.warnings.append(f"Only one student chose '{p.cfg.gender_label[g]}'. The lone member rule "
                              f"cannot be met for that student.")
        elif members:
            groups[g] = members
    return groups


def read_problem(path: Path, responses_path: Path | None, n_teams_arg: int | None) -> Problem:
    warnings.filterwarnings("ignore", message=".*extension is not supported.*")
    wb = load_workbook(path, data_only=True)
    for sheet in (lay.SETTINGS, lay.ROSTER, lay.RESPONSES, lay.OVERRIDES):
        if sheet not in wb.sheetnames:
            sys.exit(f"'{path.name}' has no '{sheet}' sheet. Is this the team formation workbook?")
    cfg = read_settings(wb)
    warn, mapping, notes, checks = [], [], [], []
    responses, dupes = read_responses(read_table(responses_path, wb), cfg, warn, mapping, checks)
    if dupes:
        notes.append(f"{dupes} earlier or repeated submission(s) replaced by the latest one.")
    roster = read_roster(wb)

    excluded: list[tuple[str, str]] = []
    if roster:
        # Microsoft Forms records the account username as the email, so there most links are by name
        students = roster
        by_email = {norm(s.email): i for i, s in enumerate(students) if s.email}
        by_words: dict[str, list[int]] = {}
        for i, s in enumerate(students):
            if name_key(s.name):
                by_words.setdefault(name_key(s.name), []).append(i)
        linked: dict[int, Student] = {}
        for r in responses:
            who = f"{r.name} <{r.email}>"
            lost = "" if is_none(r.apart_raw) else f" Its keep apart request '{r.apart_raw}' is not applied."
            i = by_email.get(norm(r.email)) if r.email else None
            if i is None:
                hits = by_words.get(name_key(r.name), []) if name_key(r.name) else []
                if len(hits) != 1:
                    why = (f"the name fits {len(hits)} roster students ({', '.join(students[j].email for j in hits)})"
                           if hits else "no roster student with this email or name")
                    excluded.append((who, f"{why}. To link it, put the student's roster email in the Email "
                                          f"cell of the response and run again.{lost}"))
                    continue
                i = hits[0]
                if norm(r.name) != norm(students[i].name):
                    checks.append(("Linked by a differently written name", f"{r.name} -> {students[i].name}"))
            prev = linked.get(i)
            if prev is not None:
                keep = r if newer(r, prev) else prev
                warn.append(f"Two responses link to {students[i].name}: {prev.name} <{prev.email}> (row {prev.row}) "
                            f"and {r.name} <{r.email}> (row {r.row}). Only the later one, row {keep.row}, is used.")
                r = keep
            linked[i] = r
        for i, r in linked.items():
            r.name, r.email = students[i].name or r.name, students[i].email or r.email
            students[i] = r
    else:
        students = responses
        warn.append("Roster is empty: only students who answered the form are placed.")
    if not students:
        sys.exit("No students found in Roster or Responses.")
    for s in students:
        s.skills = s.skills or [0] * len(cfg.skill_stmts)

    n = len(students)
    t = n_teams_arg or cfg.n_teams or math.ceil(n / cfg.max_size)
    if t > lay.TEAM_SLOTS:
        sys.exit(f"{t} teams needed but the workbook has {lay.TEAM_SLOTS} team names.")
    if t * cfg.max_size < n:
        sys.exit(f"{t} teams of at most {cfg.max_size} cannot hold {n} students.")
    if n // t < cfg.min_size:
        warn.append(f"{n} students in {t} teams gives teams of {n // t}, below the minimum of {cfg.min_size}.")
    p = Problem(path, cfg, students, t, warn, notes, mapping, [], excluded, checks)
    p.groups = rule_groups(p)

    # keep apart requests from the form, then overrides
    match = Matcher(students)
    for i, s in enumerate(students):
        apply_apart(p, match, i, s.apart_raw, s.name or s.email)
    ws = wb[lay.OVERRIDES]
    active = {norm(name): g for g, name in enumerate(cfg.team_names[:t])}
    for r in range(lay.DATA_FIRST, lay.DATA_FIRST + lay.OVERRIDE_ROWS):
        who, pin, apart = (as_str(ws.cell(r, c).value) for c in (1, 2, 3))
        if not who or is_example(who):
            continue
        i, how = match(who)
        if i is None:
            p.unmatched.append(("Overrides", who if how == "no match" else f"{who} ({how})"))
            continue
        if how not in ("email", "name"):
            p.checks.append((f"Overrides name matched by {how}", f"'{who}' -> {students[i].name}"))
        if pin:
            if norm(pin) not in active:
                sys.exit(f"Overrides row {r}: '{pin}' is not one of the {t} teams in use.")
            students[i].pin = active[norm(pin)]
        apply_apart(p, match, i, apart, f"Overrides ({who})")
    for i, s in enumerate(students):
        for j in s.apart:
            if s.pin is not None and s.pin == students[j].pin:
                sys.exit(f"{s.name} and {students[j].name} are pinned to the same team but must be kept apart.")
    return p


# ---------------------------------------------------------------------------- model
@dataclass
class Result:
    status: str                                         # OPTIMAL if every stage was proven optimal
    stages: list[tuple[str, str, int | None, float]]    # criterion, solver status, value, seconds
    wall_time: float
    terms: dict[str, int]                               # every criterion measured on the final teams


def build_model(p: Problem):
    """The hard constraints, one expression per criterion, and every variable (for hints)."""
    S, cfg, T = p.students, p.cfg, range(p.n_teams)
    N, K = range(len(S)), range(len(cfg.skill_stmts))
    lo, hi = len(S) // p.n_teams, math.ceil(len(S) / p.n_teams)
    m = cp_model.CpModel()
    made = []

    def new_bool(tag):
        made.append(m.new_bool_var(tag))
        return made[-1]

    def new_int(lo_, hi_, tag):
        made.append(m.new_int_var(lo_, hi_, tag))
        return made[-1]

    x = {(i, g): new_bool(f"x_{i}_{g}") for i in N for g in T}
    for i in N:
        m.add_exactly_one(x[i, g] for g in T)
        if S[i].pin is not None:
            m.add(x[i, S[i].pin] == 1)
    for g in T:
        m.add_linear_constraint(sum(x[i, g] for i in N), lo, hi)
    for i in N:
        for j in S[i].apart:
            if i < j:
                for g in T:
                    m.add(x[i, g] + x[j, g] <= 1)
    if all(s.pin is None for s in S):         # teams are interchangeable: break the symmetry
        for i in N:
            for g in T:
                if g > i:
                    m.add(x[i, g] == 0)

    def any_of(cands, tag):
        """At most 1 if a candidate is in the team; exact wherever the criterion is optimised or held."""
        if not cands:
            return 0
        v = new_bool(tag)
        m.add(v <= sum(cands))
        return v

    cov = {(g, k): any_of([x[i, g] for i in N if S[i].skills[k] >= cfg.threshold], f"cov{g}_{k}")
           for g in T for k in K}
    part = {(g, k): any_of([x[i, g] for i in N if S[i].skills[k] >= cfg.threshold - 1], f"part{g}_{k}")
            for g in T for k in K}
    worst = new_int(0, len(K), "worst")
    for g in T:
        m.add(worst >= len(K) - sum(cov[g, k] for k in K))

    discs = sorted({s.discipline for s in S if s.discipline})
    div = [any_of([x[i, g] for i in N if S[i].discipline == d], f"div{g}_{d}") for g in T for d in discs]

    lone = []
    for gname, members in p.groups.items():
        for g in T:
            cnt = sum(x[i, g] for i in members)
            none, two, solo = (new_bool(f"{t}_{gname}_{g}") for t in ("none", "two", "solo"))
            m.add(cnt == 0).only_enforce_if(none)
            m.add(cnt >= 2).only_enforce_if(two)
            m.add_exactly_one([none, two, solo])
            lone.append(solo)

    time_short = []
    known = [i for i in N if S[i].blocks is not None]
    if cfg.blocks and cfg.req_blocks and known:
        for g in T:
            free = []
            for b in range(len(cfg.blocks)):
                a = new_bool(f"free{g}_{b}")
                for i in known:
                    if b not in S[i].blocks:
                        m.add(a <= 1 - x[i, g])
                free.append(a)
            short = new_int(0, cfg.req_blocks, f"tshort{g}")
            m.add(short >= cfg.req_blocks - sum(free))
            time_short.append(short)

    nonresp = []
    silent = [i for i in N if not S[i].responded]
    if len(silent) > 1:
        for g in T:
            e = new_int(0, hi, f"nr{g}")
            m.add(e >= sum(x[i, g] for i in silent) - 1)
            nonresp.append(e)

    terms = {"lone": sum(lone), "time": sum(time_short), "worst": worst, "nonresp": sum(nonresp),
             "cover": sum(cov.values()), "div": sum(div), "partial": sum(part.values())}
    return m, x, terms, made


def solve(p: Problem, priority: list[str], time_limit: float, deterministic: bool, seed: int) -> Result:
    """Optimise the criteria one at a time in the order given, holding each at its best value."""
    m, x, terms, made = build_model(p)
    stages, hint, assignment, total = [], None, {}, 0.0
    for name in priority:
        sense, expr = CRITERIA[name][0], terms[name]
        (m.minimize if sense == "min" else m.maximize)(expr)
        if hint is not None:                    # start from the previous stage's teams
            m.clear_hints()
            for v, val in hint:
                m.add_hint(v, val)
        solver = cp_model.CpSolver()
        solver.parameters.random_seed = seed
        solver.parameters.num_workers = 8
        if deterministic:
            # a machine independent work budget plus interleaved search gives the same teams on
            # every run with the same input, seed and OR-Tools version
            solver.parameters.interleave_search = True
            solver.parameters.max_deterministic_time = time_limit
            solver.parameters.max_time_in_seconds = 20 * time_limit   # safety net, reported if it bites
        else:
            solver.parameters.max_time_in_seconds = time_limit
        status = solver.solve(m)
        total += solver.wall_time
        if status not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
            if hint is None and status == cp_model.INFEASIBLE:
                sys.exit("No allocation meets the hard rules (team sizes, pins, keep apart requests). "
                         "Check the pins and keep apart requests in Overrides.")
            if hint is None:
                sys.exit(f"No allocation found within the time limit (solver status "
                         f"{solver.status_name(status)}). Run again with a larger --time-limit.")
            stages.append((name, f"{solver.status_name(status)}, previous teams kept", None, solver.wall_time))
            continue
        if deterministic and status != cp_model.OPTIMAL and solver.deterministic_time < 0.99 * time_limit:
            p.warnings.append(f"Stage '{name}' stopped at the wall clock safety limit ({20 * time_limit:.0f} s) "
                              f"before its work budget was used, so this run may not be reproducible. Run again "
                              f"on a less busy computer.")
        value = int(round(solver.objective_value))
        stages.append((name, solver.status_name(status), value, solver.wall_time))
        hint = [(v, solver.value(v)) for v in made]
        assignment = {i: g for (i, g), v in x.items() if solver.boolean_value(v)}
        if not isinstance(expr, int):           # hold this criterion at its best value
            m.add(expr <= value if sense == "min" else expr >= value)
    for i, g in assignment.items():
        p.students[i].team = g
    proven = all(st == "OPTIMAL" for _, st, _, _ in stages)
    return Result("OPTIMAL" if proven else "FEASIBLE", stages, total, measure(p))


# ---------------------------------------------------------------------------- reporting
def shared_blocks(p: Problem, mem: list[Student]) -> list[int]:
    known = [s.blocks for s in mem if s.blocks is not None]
    return sorted(set.intersection(*known)) if known else list(range(len(p.cfg.blocks)))


def measure(p: Problem) -> dict[str, int]:
    """Every criterion computed from the teams themselves, not from solver variables."""
    cfg, S, T = p.cfg, p.students, range(p.n_teams)
    K = range(len(cfg.skill_stmts))
    teams = [[s for s in S if s.team == g] for g in T]
    best = [[max((s.skills[k] for s in mem), default=0) for k in K] for mem in teams]
    timed = cfg.blocks and cfg.req_blocks and any(s.blocks is not None for s in S)
    return dict(
        lone=sum(sum(S[i].team == g for i in members) == 1 for members in p.groups.values() for g in T),
        time=sum(max(0, cfg.req_blocks - len(shared_blocks(p, mem))) for mem in teams) if timed else 0,
        worst=max(sum(b < cfg.threshold for b in bt) for bt in best),
        nonresp=sum(max(0, sum(not s.responded for s in mem) - 1) for mem in teams),
        cover=sum(b >= cfg.threshold for bt in best for b in bt),
        div=sum(len({s.discipline for s in mem if s.discipline}) for mem in teams),
        partial=sum(b >= cfg.threshold - 1 for bt in best for b in bt))


def mentor(p: Problem, g: int) -> str:
    if p.cfg.mentors[g]:
        return p.cfg.mentors[g]
    return p.cfg.tas[g % len(p.cfg.tas)] if p.cfg.tas else ""


def team_rows(p: Problem) -> list[dict]:
    cfg, out = p.cfg, []
    for g in range(p.n_teams):
        mem = [s for s in p.students if s.team == g]
        best = [max((s.skills[k] for s in mem), default=0) for k in range(len(cfg.skill_stmts))]
        gaps = [cfg.skill_labels[k] for k, b in enumerate(best) if b < cfg.threshold]
        shared = shared_blocks(p, mem)
        lone = [cfg.gender_label[gk] for gk, members in p.groups.items()
                if sum(p.students[i].team == g for i in members) == 1]
        out.append(dict(
            team=cfg.team_names[g], mentor=mentor(p, g), size=len(mem),
            members=", ".join(s.name or s.email for s in mem),
            covered=len(best) - len(gaps), gaps=", ".join(gaps),
            disciplines=len({s.discipline for s in mem if s.discipline}),
            n_shared=len(shared), shared=", ".join(cfg.blocks[b] for b in shared),
            lone="lone member: " + ", ".join(lone) if lone else "OK",
            silent=", ".join(s.name or s.email for s in mem if not s.responded)))
    return out


def write_outputs(p: Problem, res: Result, priority: list[str], out_dir: Path, command: str) -> tuple[Path, Path]:
    cfg = p.cfg
    stem = p.source.stem
    xlsx_path, csv_path = out_dir / f"{stem}_teams.xlsx", out_dir / f"{stem}_teams.csv"
    head_font, head_fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    flag = PatternFill("solid", fgColor="F8CBAD")

    def table(ws, headers, rows, widths):
        for j, h in enumerate(headers, 1):
            c = ws.cell(1, j, h)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r, row in enumerate(rows, 2):
            for j, v in enumerate(row, 1):
                c = ws.cell(r, j, v)
                c.font, c.alignment = base, Alignment(wrap_text=True, vertical="top")
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + j)].width = w
        ws.freeze_panes = "B2"

    teams = team_rows(p)
    wb = Workbook()
    ws = wb.active
    ws.title = "Teams"
    table(ws, ["Team", "Mentor", "Size", "Members", f"Strong skills covered (of {len(cfg.skill_stmts)})",
               "Skills with no strong member", "Disciplines", "Shared free blocks (n)",
               "Shared free blocks", "Gender check", "No form response"],
          [[t["team"], t["mentor"], t["size"], t["members"], t["covered"], t["gaps"], t["disciplines"],
            t["n_shared"], t["shared"], t["lone"], t["silent"]] for t in teams],
          [11, 18, 6, 52, 12, 36, 11, 11, 40, 16, 28])
    for r, t in enumerate(teams, 2):
        if t["lone"] != "OK":
            ws.cell(r, 10).fill = flag
        if cfg.req_blocks and t["n_shared"] < cfg.req_blocks:
            ws.cell(r, 8).fill = flag

    ws = wb.create_sheet("Students")
    order = sorted(p.students, key=lambda s: (s.team, norm(s.name)))
    table(ws, ["Team", "Full name", "Email", "Form response", "Discipline", "Strong in", "Pinned",
               "Kept apart from"],
          [[cfg.team_names[s.team], s.name, s.email, "Yes" if s.responded else "No", s.discipline,
            ", ".join(cfg.skill_labels[k] for k, v in enumerate(s.skills) if v >= cfg.threshold),
            "Yes" if s.pin is not None else "", ", ".join(p.students[j].name for j in sorted(s.apart))]
           for s in order], [11, 28, 30, 10, 28, 44, 8, 30])

    ws = wb.create_sheet("Report")
    sizes = Counter(t["size"] for t in teams)
    pairs = p.n_teams * len(cfg.skill_stmts)
    lines = [
        ("Input workbook", str(p.source)),
        ("Run at", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Command", command),
        ("Script version", f"{__version__}; SHA-256 allocate_teams.py {sha256(Path(__file__))}, "
                           f"layout.py {sha256(Path(lay.__file__))}"),
        ("Environment", f"OR-Tools {ortools_version}, Python {platform.python_version()}, openpyxl "
                        f"{openpyxl.__version__}, {platform.system()} {platform.release()} {platform.machine()}"),
        ("", ""),
        ("Students placed", len(p.students)),
        ("With a form response", sum(s.responded for s in p.students)),
        ("Without a form response", sum(not s.responded for s in p.students)),
        ("Teams", f"{p.n_teams}: " + ", ".join(f"{n} of {k}" for k, n in sorted(sizes.items(), reverse=True))),
        ("Strong skill threshold", cfg.threshold),
        ("Shared free blocks required", cfg.req_blocks),
        ("", ""),
        ("Priority order", " > ".join(priority)),
        ("Solver", f"OR-Tools CP-SAT, one stage per criterion, {res.status}, {res.wall_time:.1f} s in total"),
    ]
    lines += [("Stage", f"{n}. {name}: {st}" + ("" if v is None else f", value {v}") + f", {secs:.1f} s")
              for n, (name, st, v, secs) in enumerate(res.stages, 1)]
    for name, (_, label) in CRITERIA.items():
        v = res.terms[name]
        lines.append((label, (f"{v} of {pairs}" if name in ("cover", "partial") else v)
                      if name in priority else f"{v} (not optimised)"))
    lines += [("", "")]
    lines += [("Column used", f"{a}: {b}") for a, b in p.mapping]
    lines += [("", "")]
    lines += [("Note", n) for n in p.notes]
    lines += [("Warning", w) for w in p.warnings]
    lines += [("Response not used", f"{who}: {why}") for who, why in p.excluded]
    lines += [(f"Check: {what}", txt) for what, txt in p.checks]
    lines += [("Keep apart request", f"{who}: {txt}") for who, txt in p.apart_log]
    lines += [("Unmatched name", f"{who}: '{txt}'") for who, txt in p.unmatched]
    lines += [
        ("", ""),
        ("Method", "Every student on the roster is assigned by a constraint programming model. Hard rules: "
                   "one team per student, team sizes differ by at most one, pins and keep apart requests "
                   "hold. The criteria are then optimised one at a time in the priority order above, each "
                   "held at its best value while the next is optimised (lexicographic optimisation): no lone "
                   "member of a gender marked Yes, the required shared free time, the worst team's skill "
                   "coverage, at most one non responder per team, total strong skill coverage, discipline "
                   "mix and total partial skill coverage. The numbers above are measured on the final teams."),
        ("Rationale", "Instructor formed, criteria based teams of 3 to 5 (Oakley, Felder, Brent & Elhajj, 2004); "
                      "no isolated women or men (Dasgupta, Scircle & Hunsinger, 2015; Kanter, 1977); shared free "
                      "time (Oakley et al., 2004); criteria based assignment as in CATME Team-Maker (Layton, "
                      "Loughry, Ohland & Ricco, 2010)."),
        ("Solver reference", "Perron, L. & Didier, F. CP-SAT, Google OR-Tools. https://developers.google.com/optimization"),
    ]
    for r, (k, v) in enumerate(lines, 1):
        a, b = ws.cell(r, 1, k), ws.cell(r, 2, v)
        a.font, b.font = bold, base
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if k in ("Warning", "Response not used"):
            b.fill = flag
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 110
    wb.move_sheet("Report", offset=-2)
    wb.active = 0
    wb.save(xlsx_path)

    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Team", "Mentor", "Full name", "Email"])
        for s in order:
            w.writerow([cfg.team_names[s.team], mentor(p, s.team), s.name, s.email])
    return xlsx_path, csv_path


# ---------------------------------------------------------------------------- CLI
def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="Form balanced student project teams from the intake workbook.")
    ap.add_argument("workbook", type=Path, help="the team formation .xlsx (Settings, Roster, Responses)")
    ap.add_argument("--responses", type=Path, help="read responses from this Forms export (.xlsx or .csv) "
                                                   "instead of the Responses sheet")
    ap.add_argument("--teams", type=int, help="number of teams (default: Settings, else fewest at max size)")
    ap.add_argument("--out-dir", type=Path, help="where to write results (default: next to the workbook)")
    ap.add_argument("--time-limit", type=float, default=30,
                    help="solver time limit per criterion in seconds (default 30); with --deterministic "
                         "a work budget per criterion")
    ap.add_argument("--deterministic", action="store_true",
                    help="reproducible run: same input, seed and OR-Tools version give the same teams")
    ap.add_argument("--seed", type=int, default=0, help="solver random seed (default 0)")
    ap.add_argument("--priority", default=",".join(DEFAULT_PRIORITY),
                    help="criteria in order of priority, comma separated; leave one out to switch it off "
                         f"(default {','.join(DEFAULT_PRIORITY)})")
    ap.add_argument("--version", action="version", version=__version__)
    args = ap.parse_args(argv)
    args.priority = [c.strip() for c in args.priority.split(",") if c.strip()]
    if not args.priority or len(set(args.priority)) < len(args.priority) or set(args.priority) - set(CRITERIA):
        ap.error(f"--priority needs distinct criteria from: {', '.join(CRITERIA)}")
    return args


def main(argv=None) -> int:
    args = parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    for path in (args.workbook, args.responses):
        if path is not None and not path.exists():
            sys.exit(f"File not found: {path}")
    p = read_problem(args.workbook, args.responses, args.teams)
    print(f"{len(p.students)} students ({sum(s.responded for s in p.students)} with a response) "
          f"into {p.n_teams} teams")
    res = solve(p, args.priority, args.time_limit, args.deterministic, args.seed)
    out_dir = args.out_dir or args.workbook.resolve().parent
    out_dir.mkdir(parents=True, exist_ok=True)
    command = "python allocate_teams.py " + " ".join(sys.argv[1:] if argv is None else argv)
    xlsx_path, csv_path = write_outputs(p, res, args.priority, out_dir, command)
    print(f"Solver: {res.status} in {res.wall_time:.1f} s | lone members {res.terms['lone']} | "
          f"missing shared blocks {res.terms['time']} | worst skill gaps {res.terms['worst']}")
    for w in p.warnings:
        print("WARNING:", w)
    for who, why in p.excluded:
        print(f"RESPONSE NOT USED: {who}: {why}")
    for who, txt in p.unmatched:
        print(f"UNMATCHED: {who}: '{txt}'")
    if p.checks:
        print(f"{len(p.checks)} match(es) to check are listed in the Report.")
    print(f"Wrote {xlsx_path}\n      {csv_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
