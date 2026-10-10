"""Build the staff intake workbook for team formation.

Staff form the teams. Students answer a private survey (Microsoft Forms or Google Forms); staff
paste the export into this workbook and run allocate_teams.py. Nothing in this workbook is shared
with students.

Usage: python build_workbook.py [output.xlsx]
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

import layout as lay

OUT = sys.argv[1] if len(sys.argv) > 1 else "team_formation.xlsx"

# ---------------------------------------------------------------- styles
FONT = "Arial"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=16, bold=True, color="1F3864")
f_h2 = Font(name=FONT, size=12, bold=True, color="1F3864")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_note = Font(name=FONT, size=9, italic=True, color="595959")
f_example = Font(name=FONT, size=10, italic=True, color="808080")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_input = PatternFill("solid", fgColor="FFF2CC")   # yellow = fill in
fill_auto = PatternFill("solid", fgColor="F2F2F2")    # grey = automatic
fill_example = PatternFill("solid", fgColor="FAFAFA")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
GREY = PatternFill(fill_type="solid", start_color="E7E6E6", end_color="E7E6E6")

wb = Workbook()
readme = wb.active
readme.title = lay.START
form = wb.create_sheet(lay.FORM)
st = wb.create_sheet(lay.SETTINGS)
ro = wb.create_sheet(lay.ROSTER)
rs = wb.create_sheet(lay.RESPONSES)
ov = wb.create_sheet(lay.OVERRIDES)
for ws, colour in ((readme, "1F3864"), (form, "2E75B6"), (st, "7F7F7F"), (ro, "70AD47"),
                   (rs, "FFC000"), (ov, "ED7D31")):
    ws.sheet_properties.tabColor = colour


def head(ws, row, col, text):
    c = ws.cell(row, col, text)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
    return c


def cell(ws, row, col, value, font=f_base, fill=None, align=None):
    c = ws.cell(row, col, value)
    c.font, c.border = font, border
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def rng(col, slots):
    """Absolute Settings range for a table column, e.g. Settings!$G$16:$G$27."""
    return (f"Settings!${L(col)}${lay.TABLE_FIRST}:${L(col)}${lay.TABLE_FIRST + slots - 1}")


B = {r: f"Settings!$B${r}" for r in range(1, 20)}

# ================================================================ Settings
st["A1"] = "Settings"
st["A1"].font = f_title
st["A2"] = ("Yellow cells are inputs. The allocator reads them. The Form questions sheet follows "
            "the question titles, statements and options entered here.")
st["A2"].font = f_note
params = [
    (lay.MAX_SIZE, "Maximum team size", 5,
     "Oakley et al. (2004) recommend 3 to 5. Above 5, free riding becomes easier."),
    (lay.MIN_SIZE, "Minimum team size", 4, "The Report warns if the cohort forces smaller teams."),
    (lay.N_TEAMS, "Number of teams (blank = automatic)", None,
     "Automatic: fewest teams at the maximum size. Sizes always differ by at most one."),
    (lay.THRESHOLD, "Strong skill threshold (1 to 5)", 4,
     "A team covers a skill if one member rates themselves at or above this value."),
    (lay.REQ_BLOCKS, "Shared free time blocks required per team", 2,
     "Blocks from column I in which every member of the team is free."),
    (lay.DEADLINE, "Form closes", "Set the date and time",
     "Shown in the form description and on the Start here sheet."),
]
for r, label, val, note in params:
    st.cell(r, 1, label).font = f_bold
    cell(st, r, 2, val, fill=fill_input, align=Alignment(horizontal="center"))
    st.cell(r, 3, note).font = f_note
roster_rng = f"Roster!$B${lay.DATA_FIRST}:$B${lay.DATA_FIRST + lay.ROSTER_ROWS - 1}"
resp_rng = "Responses!$A$2:$A$2001"
autos = [
    (lay.ROSTER_COUNT, "Students on roster (auto)",
     f'=COUNTIF({roster_rng},"?*")', "Emails in the Roster sheet."),
    (lay.RESPONSE_COUNT, "Responses pasted (auto)",
     f'=COUNTIF({resp_rng},"?*")-COUNTIF({resp_rng},"Example*")',
     "Rows in the Responses sheet, without the example row."),
    (lay.TEAMS_EFFECTIVE, "Teams (auto)",
     f'=IF({B[lay.N_TEAMS]}<>"",{B[lay.N_TEAMS]},IF({B[lay.ROSTER_COUNT]}=0,"",'
     f'ROUNDUP({B[lay.ROSTER_COUNT]}/{B[lay.MAX_SIZE]},0)))',
     "Uses the roster size. The allocator recomputes this from the data it reads."),
    (lay.SIZES_TEXT, "Team sizes (auto)",
     f'=IF(OR({B[lay.TEAMS_EFFECTIVE]}="",{B[lay.ROSTER_COUNT]}=0),"",'
     f'MOD({B[lay.ROSTER_COUNT]},{B[lay.TEAMS_EFFECTIVE]})&" of "&'
     f'(INT({B[lay.ROSTER_COUNT]}/{B[lay.TEAMS_EFFECTIVE]})+1)&", "&'
     f'({B[lay.TEAMS_EFFECTIVE]}-MOD({B[lay.ROSTER_COUNT]},{B[lay.TEAMS_EFFECTIVE]}))&" of "&'
     f'INT({B[lay.ROSTER_COUNT]}/{B[lay.TEAMS_EFFECTIVE]}))',
     "Number of teams of each size, for example 11 of 5, 1 of 4."),
]
for r, label, formula, note in autos:
    st.cell(r, 1, label).font = f_bold
    cell(st, r, 2, formula, fill=fill_auto, align=Alignment(horizontal="center"))
    st.cell(r, 3, note).font = f_note
for r in range(3, lay.TABLE_HEAD):
    st.cell(r, 3).alignment = Alignment(wrap_text=False)   # notes overflow into empty columns

H = lay.TABLE_HEAD
for col, text in ((lay.TEAM_COL, "Team names"), (lay.MENTOR_COL, "Mentor (blank = rotate mentor list)"),
                  (lay.SKILL_LABEL_COL, "Skill: short label"),
                  (lay.SKILL_STMT_COL, "Skill: statement in the form"),
                  (lay.DISC_COL, "Discipline options"), (lay.BLOCK_COL, "Free time blocks"),
                  (lay.GENDER_COL, "Gender options"),
                  (lay.GENDER_RULE_COL, "Avoid a lone member? (Yes/No)"), (lay.TA_COL, "Mentor list"),
                  (lay.KEY_LABEL_COL, "Column"),
                  (lay.KEY_TEXT_COL, "Header text the allocator looks for (= form question title)")):
    head(st, H, col, text)
st.row_dimensions[H].height = 42


def fill_column(col, values, slots, align=None):
    for i in range(slots):
        cell(st, lay.TABLE_FIRST + i, col, values[i] if i < len(values) else None,
             fill=fill_input, align=align)


fill_column(lay.TEAM_COL, [f"Team {g + 1:02d}" for g in range(lay.TEAM_SLOTS)], lay.TEAM_SLOTS)
fill_column(lay.MENTOR_COL, [], lay.TEAM_SLOTS)
fill_column(lay.SKILL_LABEL_COL, [s for s, _ in lay.SKILLS], lay.SKILL_SLOTS)
fill_column(lay.SKILL_STMT_COL, [d for _, d in lay.SKILLS], lay.SKILL_SLOTS)
fill_column(lay.DISC_COL, lay.DISCIPLINES, lay.DISC_SLOTS)
fill_column(lay.BLOCK_COL, lay.BLOCKS, lay.BLOCK_SLOTS)
fill_column(lay.GENDER_COL, [g for g, _ in lay.GENDERS], lay.GENDER_SLOTS)
fill_column(lay.GENDER_RULE_COL, ["Yes" if r else "No" for _, r in lay.GENDERS], lay.GENDER_SLOTS,
            align=Alignment(horizontal="center"))
fill_column(lay.TA_COL, lay.TAS, lay.TA_SLOTS)
for i, (key, text, note) in enumerate(lay.KEYS):
    r = lay.TABLE_FIRST + i
    c = cell(st, r, lay.KEY_LABEL_COL, key, font=f_bold, fill=fill_auto)
    c.comment = Comment(note, "Module team")
    cell(st, r, lay.KEY_TEXT_COL, text, fill=fill_input, align=wrap_top)

widths = {lay.TEAM_COL: 30, lay.MENTOR_COL: 22, 3: 3, lay.SKILL_LABEL_COL: 18,
          lay.SKILL_STMT_COL: 34, 6: 3, lay.DISC_COL: 30, 8: 3, lay.BLOCK_COL: 20, 10: 3,
          lay.GENDER_COL: 34, lay.GENDER_RULE_COL: 14, 13: 3, lay.TA_COL: 22, 15: 3,
          lay.KEY_LABEL_COL: 14, lay.KEY_TEXT_COL: 58}
for col, w in widths.items():
    st.column_dimensions[L(col)].width = w
st.freeze_panes = f"A{lay.TABLE_FIRST}"

for r, lo, hi in ((lay.MAX_SIZE, 2, 8), (lay.MIN_SIZE, 1, 8), (lay.N_TEAMS, 1, lay.TEAM_SLOTS),
                  (lay.THRESHOLD, 1, 5), (lay.REQ_BLOCKS, 0, lay.BLOCK_SLOTS)):
    dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi),
                        allow_blank=(r == lay.N_TEAMS), showErrorMessage=True,
                        error=f"Enter a whole number from {lo} to {hi}.")
    st.add_data_validation(dv)
    dv.add(f"B{r}")
team_first, team_last = lay.TABLE_FIRST, lay.TABLE_FIRST + lay.TEAM_SLOTS - 1
ta_col = L(lay.TA_COL)
ta_cells = f"${ta_col}${lay.TABLE_FIRST}:${ta_col}${lay.TABLE_FIRST + lay.TA_SLOTS - 1}"
dv = DataValidation(type="list", allow_blank=True,
                    formula1=f"OFFSET(${ta_col}${lay.TABLE_FIRST},0,0,MAX(1,COUNTA({ta_cells})),1)")
st.add_data_validation(dv)
dv.add(f"B{team_first}:B{team_last}")
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
st.add_data_validation(dv)
dv.add(f"L{lay.TABLE_FIRST}:L{lay.TABLE_FIRST + lay.GENDER_SLOTS - 1}")
st.conditional_formatting.add(
    f"A{team_first}:B{team_last}",
    FormulaRule(formula=[f'AND(ISNUMBER($B${lay.TEAMS_EFFECTIVE}),ROW()-{team_first - 1}>$B${lay.TEAMS_EFFECTIVE})'],
                fill=GREY, font=Font(color="A6A6A6")))

# ================================================================ Form questions
KEYROW = {key: lay.TABLE_FIRST + i for i, (key, _, _) in enumerate(lay.KEYS)}
Q = {key: f"Settings!$Q${row}" for key, row in KEYROW.items()}
form["A1"] = "Survey to create in Microsoft Forms or Google Forms (students answer privately)"
form["A1"].font = f_title
form["A2"] = ("Create a survey with the questions below, in this order. Copy titles exactly: the "
              "allocator finds columns by these titles. Grey cells update from Settings.")
form["A2"].font = f_note
form["A4"] = "Form settings"
form["A4"].font = f_h2
form["B5"] = ("Microsoft Forms: only people in my organisation can respond; tick Record name and One response "
              "per person. Google Forms: restrict to your organisation, collect verified email addresses, limit "
              "to one response, and set the completion time column title in Settings column Q to Timestamp.")
form["A7"] = "Form description"
form["A7"].font = f_h2
form["B8"] = (
    '="Staff form the project teams for the group project. We use your '
    'answers for three things. First, we spread skills and disciplinary backgrounds across teams. '
    'Second, every team shares at least "&Settings!$B$7&" free time blocks outside class. Third, we '
    'avoid teams in which one person is the only woman or the only man. Your answers are visible to '
    'the module staff only. They are not used for marking, and we delete them once the teams are '
    'final. The gender question is optional. If anything else should be taken into account, for '
    'example a disability or caring responsibilities, contact the teaching team privately. The form '
    'closes "&Settings!$B$8&"."')
for ref in ("B5", "B8"):
    form[ref].font = f_base
    form[ref].alignment = wrap_top
    form[ref].fill = fill_auto
form.row_dimensions[5].height = 30
form.row_dimensions[8].height = 95

qhead = ["#", "Question title (copy exactly)", "Forms question type", "Options or statements",
         "Required", "Note"]
QH = 10
for j, h in enumerate(qhead, 1):
    head(form, QH, j, h)
disc = rng(lay.DISC_COL, lay.DISC_SLOTS)
blocks = rng(lay.BLOCK_COL, lay.BLOCK_SLOTS)
genders = rng(lay.GENDER_COL, lay.GENDER_SLOTS)
questions = [
    ("auto", f"={Q['name']}&\" and \"&{Q['email']}", "Recorded automatically",
     "No question needed", "", "Microsoft Forms: tick Record name. Google Forms: collect email addresses."),
    (1, f"={Q['discipline']}", "Choice (one answer)", f'=_xlfn.TEXTJOIN("; ",TRUE,{disc})',
     "Yes", "Used to mix disciplines within teams."),
    (2, f"={Q['skills']}", "Likert",
     f'="Options: {"; ".join(lay.RATING_OPTIONS)}. Statements: see rows below."', "Yes",
     "Self rating. Used only to balance teams, never for marking."),
]
r = QH + 1
for num, title, qtype, opts, req, note in questions:
    vals = [num, title, qtype, opts, req, note]
    for j, v in enumerate(vals, 1):
        cell(form, r, j, v, fill=fill_auto if j in (2, 4) else None, align=wrap_top)
    r += 1
for i in range(lay.SKILL_SLOTS):
    sref = f"Settings!${L(lay.SKILL_STMT_COL)}${lay.TABLE_FIRST + i}"
    cell(form, r, 1, None)
    cell(form, r, 2, None)
    cell(form, r, 3, f'=IF({sref}="","","Likert statement")', font=f_note)
    cell(form, r, 4, f'=IF({sref}="","",{sref})', fill=fill_auto)
    cell(form, r, 5, None)
    cell(form, r, 6, None)
    r += 1
more = [
    (3, f"={Q['availability']}", "Choice (multiple answers)",
     f'=_xlfn.TEXTJOIN("; ",TRUE,{blocks})', "Yes",
     "Edit the blocks in Settings column I. Do not use semicolons or commas in labels."),
    (4, f"={Q['gender']}", "Choice (one answer)", f'=_xlfn.TEXTJOIN("; ",TRUE,{genders})', "No",
     "Used only for the lone member rule in Settings column L."),
    (5, f"={Q['apart']}", "Text (long answer)", "Full names, separated by semicolons", "No",
     "Safeguarding. Always respected. Visible to staff only."),
]
for num, title, qtype, opts, req, note in more:
    for j, v in enumerate([num, title, qtype, opts, req, note], 1):
        cell(form, r, j, v, fill=fill_auto if j in (2, 4) else None, align=wrap_top)
    r += 1
for col, w in {"A": 6, "B": 52, "C": 26, "D": 60, "E": 10, "F": 50}.items():
    form.column_dimensions[col].width = w
for rr in range(QH + 1, r):
    form.row_dimensions[rr].height = 30
form.merge_cells("B5:F5")
form.merge_cells("B8:F8")
form.sheet_view.showGridLines = False

# ================================================================ Roster
ro["A1"], ro["B1"] = "Full name", "Email"
for j in (1, 2):
    head(ro, 1, j, ro.cell(1, j).value)
ro.cell(1, 3, "Paste the class list from your learning platform (students only: full name and email) "
              "from row 3. Email and name link each student to their survey response.").font = f_note
cell(ro, 2, 1, "Example: Alex Taylor", font=f_example, fill=fill_example)
cell(ro, 2, 2, "alex.taylor@example.edu", font=f_example, fill=fill_example)
for rr in range(lay.DATA_FIRST, lay.DATA_FIRST + lay.ROSTER_ROWS):
    for j in (1, 2):
        cell(ro, rr, j, None, fill=fill_input)
ro.column_dimensions["A"].width = 32
ro.column_dimensions["B"].width = 34
ro.freeze_panes = "A3"

# ================================================================ Responses
resp_head = (["Id", "Start time", "Completion time", lay.KEYS[1][1], lay.KEYS[0][1], lay.KEYS[3][1]]
             + [f"{lay.KEYS[4][1]}.{d}" for _, d in lay.SKILLS]
             + [lay.KEYS[5][1], lay.KEYS[6][1], lay.KEYS[7][1]])
example = (["Example", "10/9/26 14:02", "10/9/26 14:06", "alex.taylor@example.edu",
            "Example: Alex Taylor", "Psychology / cognitive science"]
           + ["4", "2", "3 done it in a project", "3", "1 no experience", "1", "4", "2", "4"]
           + ["Monday morning;Wednesday afternoon;Thursday evening;", "Prefer not to say", ""])
for j, h in enumerate(resp_head, 1):
    head(rs, 1, j, h)
    cell(rs, 2, j, example[j - 1], font=f_example, fill=fill_example)
    rs.column_dimensions[L(j)].width = 18
rs.row_dimensions[1].height = 60
rs.cell(1, len(resp_head) + 2, "Paste the survey export here, starting in A1 (headers included): "
                               "Microsoft Forms, Open results in Excel; Google Forms, Download responses "
                               "(.csv). This replaces the example. Column order does not matter."
        ).font = f_note
rs.freeze_panes = "A2"

# ================================================================ Overrides
ov_head = ["Student (name or email)", "Pin to team (team name)",
           "Keep apart from (names or emails; separate with ;)", "Reason (staff note)"]
for j, h in enumerate(ov_head, 1):
    head(ov, 1, j, h)
ov_ex = ["Example: Alex Taylor", "Team 03", "Sam Patel", "Reasonable adjustment (agreed with disability services)"]
for j, v in enumerate(ov_ex, 1):
    cell(ov, 2, j, v, font=f_example, fill=fill_example)
for rr in range(lay.DATA_FIRST, lay.DATA_FIRST + lay.OVERRIDE_ROWS):
    for j in range(1, 5):
        cell(ov, rr, j, None, fill=fill_input)
for col, w in {"A": 30, "B": 22, "C": 44, "D": 44}.items():
    ov.column_dimensions[col].width = w
ov.row_dimensions[1].height = 32
ov.freeze_panes = "A3"
dv = DataValidation(type="list", allow_blank=True, showErrorMessage=True,
                    formula1=f"OFFSET(Settings!$A${lay.TABLE_FIRST},0,0,MAX(1,Settings!$B${lay.TEAMS_EFFECTIVE}),1)",
                    error="Choose a team name from the list.")
ov.add_data_validation(dv)
dv.add(f"B{lay.DATA_FIRST}:B{lay.DATA_FIRST + lay.OVERRIDE_ROWS - 1}")

# ================================================================ Start here
readme.column_dimensions["A"].width = 3
readme.column_dimensions["B"].width = 112
rows = [
    ("title", "Team formation (staff only)"),
    ("formula", f'="Form closes: "&{B[lay.DEADLINE]}&"   ·   Teams: "&{B[lay.TEAMS_EFFECTIVE]}&'
                f'"   ·   Sizes: "&{B[lay.SIZES_TEXT]}&"   ·   Responses: "&{B[lay.RESPONSE_COUNT]}&'
                f'" of "&{B[lay.ROSTER_COUNT]}'),
    ("gap", None),
    ("p", "Staff form the teams from a private survey. Students do not choose their team and do not see "
          "this workbook. It contains personal data, so keep it in a staff only location."),
    ("gap", None),
    ("h2", "Steps"),
    ("p", "1.  Settings: check team size, skill threshold, required shared time blocks, the skill, "
          "discipline, time block and gender lists, the mentor list and the deadline."),
    ("p", "2.  Form questions: create the survey exactly as listed and post its link on your learning platform. "
          "The form description tells students how their answers are used."),
    ("p", "3.  After the deadline, paste the class list (students only: full name and email) into Roster from "
          "row 3. Download the responses (Microsoft Forms: Open results in Excel; Google Forms: Download "
          "responses), copy everything and paste it into Responses at A1."),
    ("p", "4.  Optional: Overrides. Pin a student to a team (for example after a reasonable adjustment "
          "agreed with disability services) or keep two students apart."),
    ("p", "5.  Run the allocator: drag this file onto allocate.bat (Windows) or run ./allocate.sh FILE "
          "(macOS, Linux). It writes <name>_teams.xlsx and <name>_teams.csv next to the file."),
    ("p", "6.  Check the Report sheet of the output, especially warnings and unmatched names. Then "
          "create the groups in your learning platform from the CSV and announce the teams with their mentors."),
    ("p", "7.  Once the teams are final, delete the Responses data and the output files."),
    ("gap", None),
    ("h2", "What the allocator does"),
    ("p", "Hard rules: every student on the roster is placed once; team sizes differ by at most one; "
          "pins and keep apart requests are always respected. Students on the roster without a response "
          "are placed too, at most one per team where possible."),
    ("p", "Priorities, in strict order (a lower priority never overrides a higher one): (1) no team has "
          "exactly one woman or exactly one man; (2) every team shares the required free time blocks; (3) the "
          "team with the most skill gaps has as few as possible; (4) at most one non responder per team; (5) "
          "a strong member for each skill in each team; (6) a mix of disciplines; (7) members close to the "
          "skill threshold. The Report lists the result for each priority."),
    ("p", "The lone member rule applies only to the gender options marked Yes in Settings. It is off for "
          "non-binary and self described answers, because it would force those students into the same "
          "team and could disclose their gender to teammates."),
    ("gap", None),
    ("h2", "After the teams are formed"),
    ("p", "In the first weeks, each team writes a team charter (roles, communication channel, internal "
          "deadlines, consequence of missed work) and discusses it with its mentor (Mathieu & Rapp, 2009)."),
    ("p", "Give the teams a small, low stakes first task early in the term, so dysfunctional teams show "
          "up while there is time to intervene."),
    ("p", "Peer assessment of contribution (Loughry, Ohland & Moore, 2007) reduces free riding. Before "
          "announcing peer adjusted marks, check that your institution's assessment regulations allow them. "
          "mod_peerwork (open source Moodle plugin) and WebPA implement the adjustment."),
    ("gap", None),
    ("h2", "Evidence"),
    ("p", "Instructor formed teams based on explicit criteria work better than self selected teams when "
          "students do not know each other. Self selection produces homogeneous teams and leaves a group "
          "of students nobody chose (Oakley, Felder, Brent & Elhajj, 2004). Positive reports on self "
          "selection come from cohorts who already knew each other (Bacon, Stewart & Silver, 1999)."),
    ("p", "Women in female minority engineering teams participate less and report more anxiety than women "
          "in female majority teams (Dasgupta, Scircle & Hunsinger, 2015). Numerically isolated members "
          "in general are treated as tokens (Kanter, 1977). The allocator therefore avoids lone members."),
    ("p", "Lack of shared free time is a frequent practical cause of team breakdown (Oakley et al., 2004). "
          "Heterogeneous grouping benefits low achieving students and does not harm high achieving "
          "students (Lou et al., 1996). The survey collects no nationality or first language: cultural "
          "diversity raises creativity and lowers cohesion (Stahl, Maznevski, Voigt & Jonsen, 2010), so "
          "it offers no clear target, and discipline already breaks up most same background clusters."),
    ("p", "Telling students the criteria makes the process transparent and reduces complaints about "
          "fairness (Michaelsen & Sweet, 2008). Criteria based allocation follows CATME Team-Maker "
          "(Layton, Loughry, Ohland & Ricco, 2010)."),
    ("ref", "Bacon, D. R., Stewart, K. A., & Silver, W. S. (1999). Lessons from the best and worst student "
            "team experiences. Journal of Management Education, 23(5), 467 to 488."),
    ("ref", "Dasgupta, N., Scircle, M. M., & Hunsinger, M. (2015). Female peers in small work groups enhance "
            "women's motivation, verbal participation, and career aspirations in engineering. PNAS, 112(16), "
            "4988 to 4993."),
    ("ref", "Kanter, R. M. (1977). Some effects of proportions on group life. American Journal of Sociology, "
            "82(5), 965 to 990."),
    ("ref", "Layton, R. A., Loughry, M. L., Ohland, M. W., & Ricco, G. D. (2010). Design and validation of a "
            "web-based system for assigning members to teams using instructor-specified criteria. Advances "
            "in Engineering Education, 2(1), 1 to 28."),
    ("ref", "Lou, Y., Abrami, P. C., Spence, J. C., Poulsen, C., Chambers, B., & d'Apollonia, S. (1996). "
            "Within-class grouping: A meta-analysis. Review of Educational Research, 66(4), 423 to 458."),
    ("ref", "Loughry, M. L., Ohland, M. W., & Moore, D. D. (2007). Development of a theory-based assessment "
            "of team member effectiveness. Educational and Psychological Measurement, 67(3), 505 to 524."),
    ("ref", "Mathieu, J. E., & Rapp, T. L. (2009). Laying the foundation for successful team performance "
            "trajectories: The roles of team charters and performance strategies. Journal of Applied "
            "Psychology, 94(1), 90 to 103."),
    ("ref", "Michaelsen, L. K., & Sweet, M. (2008). The essential elements of team-based learning. New "
            "Directions for Teaching and Learning, 2008(116), 7 to 27."),
    ("ref", "Oakley, B., Felder, R. M., Brent, R., & Elhajj, I. (2004). Turning student groups into effective "
            "teams. Journal of Student Centered Learning, 2(1), 9 to 34."),
    ("ref", "Stahl, G. K., Maznevski, M. L., Voigt, A., & Jonsen, K. (2010). Unraveling the effects of cultural "
            "diversity in teams: A meta-analysis of research on multicultural work groups. Journal of "
            "International Business Studies, 41(4), 690 to 709."),
]
r = 1
for kind, text in rows:
    c = readme.cell(r, 2, text)
    c.alignment = wrap_top
    if kind == "title":
        c.font = f_title
        readme.row_dimensions[r].height = 26
    elif kind == "formula":
        c.font = Font(name=FONT, size=11, bold=True, color="C00000")
    elif kind == "h2":
        c.font = f_h2
        readme.row_dimensions[r].height = 20
    elif kind == "p":
        c.font = f_base
        readme.row_dimensions[r].height = 13.5 * max(1, -(-len(text) // 130))
    elif kind == "ref":
        c.font = f_note
        readme.row_dimensions[r].height = 12.5 * max(1, -(-len(text) // 150))
    r += 1
readme.sheet_view.showGridLines = False

wb.active = 0
wb.save(OUT)
print("saved", OUT)
